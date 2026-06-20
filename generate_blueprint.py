import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_blueprint():
    doc = docx.Document()
    
    # Title
    title = doc.add_heading('Project Blueprint: Prompt Injection Attack Simulator & Webpage Scanner', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('Document Date: June 14, 2026')
    doc.add_paragraph('Status: Blueprint Generated from PRD & SRS')
    
    doc.add_page_break()

    # 1. Executive Summary
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph(
        'The Prompt Injection Attack Simulator & Webpage Scanner is an educational and defensive '
        'security web application designed to help developers, security testers, and researchers '
        'identify potential indirect prompt-injection vulnerabilities in web content. It accepts a '
        'URL or HTML file, scans visible and hidden text-bearing sources, detects suspicious '
        'natural-language instructions targeting AI models, calculates a risk score, and generates '
        'defensive reports without executing the page content.'
    )

    # 2. System Architecture
    doc.add_heading('2. System Architecture', level=1)
    doc.add_paragraph('The system is built on a client-server architecture with the following main components:')
    
    arch_bullets = [
        'React Frontend: A single-page application built with React, React Router, and Tailwind CSS '
        'handling user input, displaying progress, rendering evidence, and requesting reports.',
        'FastAPI Backend: A Python-based REST API handling input validation, secure target fetching, '
        'HTML parsing, and analysis orchestration.',
        'Secure URL Fetcher: A specialized module ensuring SSRF mitigation by validating URLs, '
        'rejecting internal IP addresses, and safely streaming responses.',
        'HTML Parser and Extractors: Uses BeautifulSoup to extract various sources including '
        'comments, metadata, ARIA labels, image alt text, and hidden/visible text without execution.',
        'Detection Engine: A rule-based engine utilizing regex and phrase matching to identify '
        'suspicious instruction overrides, exfiltration attempts, and role manipulations.',
        'Risk & Recommendation Engines: Calculates a 0-100 risk score and maps findings to actionable mitigations.',
        'Report Generator: Creates self-contained HTML (and later PDF) reports.'
    ]
    for bullet in arch_bullets:
        doc.add_paragraph(bullet, style='List Bullet')

    # 3. User Personas & Journeys
    doc.add_heading('3. User Personas & Journeys', level=1)
    doc.add_heading('Target Personas', level=2)
    personas = [
        'Persona A (Security Student): Focuses on learning through safe demo labs to understand hiding techniques.',
        'Persona B (AI Application Developer): Focuses on validating webpage sanitization controls and reading mitigation guidance.',
        'Persona C (Security Analyst): Focuses on consistent scoring, evidence extraction, and generating exportable reports.'
    ]
    for persona in personas:
        doc.add_paragraph(persona, style='List Bullet')
        
    doc.add_heading('Primary Scan Journey', level=2)
    journey = [
        'User opens Scanner and selects URL or HTML File.',
        'User inputs target and starts scan.',
        'Backend validates target, safely retrieves HTML, and parses it.',
        'Backend extracts content and applies detection rules.',
        'Backend calculates risk score and formats findings.',
        'Results page displays findings, severities, evidence snippets, and recommendations.',
        'User can optionally export an HTML report.'
    ]
    for step in journey:
        doc.add_paragraph(step, style='List Number')

    # 4. Core Features & Functional Requirements
    doc.add_heading('4. Core Features & Functional Requirements', level=1)
    features = [
        ('Input Modes', 'Support for HTTP/HTTPS URLs and local HTML file uploads.'),
        ('Content Extraction', 'Extracts from Comments, Meta tags, Alt attributes, ARIA labels, Title, '
         'Data attributes, Inline CSS/hidden elements, and visible text.'),
        ('Rule-Based Detection', 'Detects Instruction Overrides, Data Exfiltration, Unauthorized Actions, '
         'Role Manipulation, and Stealth Instructions.'),
        ('Scoring System', '0-100 deterministic risk score categorized into Low (0-20), Medium (21-50), '
         'High (51-80), and Critical (81-100).'),
        ('Demo Lab', 'Bundled safe examples demonstrating comment injection, CSS injection, etc.'),
        ('Reporting', 'Exportable HTML report generation summarizing findings, severity, and mitigations.')
    ]
    for feature, desc in features:
        p = doc.add_paragraph()
        p.add_run(feature + ': ').bold = True
        p.add_run(desc)

    # 5. Security & Privacy Model
    doc.add_heading('5. Security & Privacy Model', level=1)
    sec_bullets = [
        'SSRF Prevention: Strict target validation, blocking localhost, private ranges, cloud metadata, '
        'and enforcing safe redirects.',
        'No Execution: Static HTML scanning only; JavaScript and active payloads are completely ignored.',
        'Output Encoding: All dynamic content and evidence strings are safely escaped before UI rendering '
        'to prevent XSS.',
        'Privacy: Raw HTML and uploads are transient and not persisted by default. Reports do not contain '
        'raw page content or authorization secrets.'
    ]
    for bullet in sec_bullets:
        doc.add_paragraph(bullet, style='List Bullet')

    # 6. Technical Stack & Deployment
    doc.add_heading('6. Technical Stack & Deployment', level=1)
    stack = [
        ('Frontend', 'React, React Router, Tailwind CSS, Axios/Fetch.'),
        ('Backend', 'Python 3.11+, FastAPI, Pydantic, BeautifulSoup4, httpx.'),
        ('Testing', 'pytest, Vitest, React Testing Library.'),
        ('Deployment Constraints', 'HTTPS required, backend to run as non-root, restrictive CORS, '
         'and comprehensive rate-limiting.'),
        ('Database (Optional Phase 3)', 'PostgreSQL for robust history, SQLite for local.')
    ]
    for item, desc in stack:
        p = doc.add_paragraph()
        p.add_run(item + ': ').bold = True
        p.add_run(desc)

    # 7. Release Plan
    doc.add_heading('7. Release Plan', level=1)
    releases = [
        ('Release 1: Core Engine', 'File uploads, static parsing, detection rules, and risk scoring logic.'),
        ('Release 2: Web MVP', 'URL scanner with SSRF controls, React UI, Results, Demo Labs, HTML Report export.'),
        ('Release 3: Product Completion', 'Dashboard, scan history, PDF export, deployment hardening.'),
        ('Release 4: Advanced Analysis', 'Sandboxed DOM rendering (Playwright), AI classification, CI/CD integrations.')
    ]
    for rel, desc in releases:
        p = doc.add_paragraph()
        p.add_run(rel + ': ').bold = True
        p.add_run(desc)

    # Save Document
    filename = 'ShadowPrompt_Blueprint.docx'
    doc.save(filename)
    print(f'Successfully generated {filename}')

if __name__ == '__main__':
    create_blueprint()
